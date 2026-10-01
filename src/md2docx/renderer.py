"""Markdown AST → docx 渲染引擎。

把 markdown-it 的 token 流翻译成 Word 段落/表格/图片，同时按配置套用排版参数。
两条原则：
1. 每段/每个 run 都写"显式格式化"，不依赖样式继承。这样在 Word 和 WPS 里
   表现一致，也不会被主题字体悄悄改掉。
2. markdown 的语法标记一律不进正文，只有内容进正文。
"""

from __future__ import annotations

import os
import re
import io
import json
from urllib.parse import unquote
from dataclasses import dataclass, replace

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.run import Run
from markdown_it import MarkdownIt
from docbridge_math import FormulaError, latex_to_omml, render_png
from docbridge_math.markdown import install_math_rules

from . import ooxml
from .config import element_style, heading_numbering
from .mermaid import MermaidRenderer
from .numbering import (BULLET_NUM_ID, HEADING_NUM_ID, ORDERED_NUM_ID,
                        ensure_numbering_part, heading_plan, list_plans)
from .units import (parse_chars, parse_length_pt, parse_size,
                    strip_heading_prefix)

STYLE_NAMES = {
    "body": "Normal",
    "heading1": "Heading 1", "heading2": "Heading 2", "heading3": "Heading 3",
    "heading4": "Heading 4", "heading5": "Heading 5", "heading6": "Heading 6",
    "list": "MD List",
    "list_ordered": "MD List Ordered",
    "quote": "MD Quote",
    "code": "MD Code",
    "table_text": "MD Table Text",
    "table_header": "MD Table Header",
    "caption": "MD Caption",
    "hr": "MD HR",
    "image": "MD Image",
}

CUSTOM_STYLES = ("MD List", "MD List Ordered", "MD Quote", "MD Code",
                 "MD Table Text", "MD Table Header", "MD Caption", "MD HR",
                 "MD Image")

HEADING_CLOSE = {"strong_close", "em_close", "s_close", "link_close"}
HEADING_OPEN = {"strong_open", "em_open", "s_open", "link_open"}

ALIGN_ENUM = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def is_cjk(ch: str) -> bool:
    if not ch:
        return False
    code = ord(ch)
    return (0x2E80 <= code <= 0x303F or 0x3040 <= code <= 0x33FF
            or 0x3400 <= code <= 0x4DBF or 0x4E00 <= code <= 0x9FFF
            or 0xF900 <= code <= 0xFAFF or 0xFF00 <= code <= 0xFFEF
            or 0x20000 <= code <= 0x2FA1F)


@dataclass
class RunCtx:
    """一次 run 的字符格式上下文（强调/链接/等宽的叠加状态）。"""
    bold: bool | None = None
    italic: bool | None = None
    strike: bool | None = None
    underline: bool | None = None
    mono: bool = False
    link: str | None = None


class MarkdownToDocx:
    def __init__(self, config: dict, md_path: str, out_path: str, log=print):
        self.cfg = config
        self.md_path = md_path
        self.out_path = os.path.abspath(out_path)
        self.base_dir = os.path.dirname(os.path.abspath(md_path))
        self.log = log
        self.warnings: list[str] = []
        self.counters = {"paragraph": 0, "heading": 0, "table": 0, "list": 0,
                         "code": 0, "mermaid": 0, "image": 0, "formula": 0}

        self.md = MarkdownIt("commonmark", {"html": False, "linkify": False})
        self.md.enable(["table", "strikethrough"])
        install_math_rules(self.md)

        self.doc = Document()
        self.text_width_emu = 0
        # 标题自动编号的配置（未开启时为 None）
        self._heading_numbering = heading_numbering(config)
        self._setup_page()
        self._setup_styles()
        self._mermaid = None
        self._cache_dir = os.path.join(
            os.path.dirname(self.out_path),
            str(config.get("mermaid", {}).get("cache_dir", "md2docx-assets")),
            os.path.splitext(os.path.basename(md_path))[0])

    # ------------------------------------------------------------------ 初始化
    def _setup_page(self) -> None:
        from .ooxml import apply_page

        page_cfg = self.cfg.get("page") or {}
        section = self.doc.sections[0]
        apply_page(section, page_cfg)
        self.section = section
        self.text_width_emu = int(section.page_width - section.left_margin
                                  - section.right_margin)
        if page_cfg.get("footer_page_number"):
            footer = section.footer
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.alignment = ALIGN_ENUM["center"]
            ooxml.add_page_number_field(para)

    def _setup_styles(self) -> None:
        ooxml.apply_doc_defaults(self.doc, element_style(self.cfg, "body"))
        for key, name in STYLE_NAMES.items():
            if key.startswith("heading") and name not in [s.name for s in self.doc.styles]:
                continue
            style = self._ensure_style(name)
            self._format_style(style, key)

    def _ensure_style(self, name: str):
        try:
            return self.doc.styles[name]
        except KeyError:
            style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            try:
                style.base_style = self.doc.styles["Normal"]
            except KeyError:
                pass
            return style

    def _format_style(self, style, key: str) -> None:
        s = element_style(self.cfg, key)
        size_pt = parse_size(s.get("size", "小四"))
        ooxml.apply_run_format(ooxml.style_rpr(style),
                               font_zh=s.get("font_zh"), font_en=s.get("font_en"),
                               size_pt=size_pt,
                               bold=s.get("bold"), italic=s.get("italic"),
                               color=s.get("color"))
        ooxml.apply_para_format(
            style.element.get_or_add_pPr(),
            line_spacing=s.get("line_spacing", 1.25),
            space_before_pt=parse_length_pt(s.get("space_before", 0)),
            space_after_pt=parse_length_pt(s.get("space_after", 0)),
            align=s.get("align"), first_line=s.get("first_line_indent"),
            left=s.get("left_indent"), hanging=s.get("hanging_indent"),
            keep_next=s.get("keep_with_next"), keep_lines=s.get("keep_lines"),
            page_break_before=s.get("page_break_before"),
            shading=s.get("shading"), borders=s.get("borders"),
            font_pt_for_chars=size_pt,
            contextual_spacing=s.get("contextual_spacing"))
        if key.startswith("heading"):
            level = int(key[-1])
            ooxml.apply_para_format(style.element.get_or_add_pPr(),
                                    outline_level=level - 1)
            nxt = OxmlElement("w:next")
            nxt.set(qn("w:val"), "Normal")
            existing = style.element.find(qn("w:next"))
            if existing is not None:
                style.element.replace(existing, nxt)
            else:
                # w:next 必须排在 pPr/rPr 之前，直接用 schema 顺序插入
                ooxml.insert_ordered(style.element, nxt, ooxml.STYLE_ORDER)

    # ------------------------------------------------------------------ 工具
    def _params(self, key: str) -> tuple[dict, dict]:
        """返回 (段落参数, 字符参数)。"""
        s = element_style(self.cfg, key)
        size_pt = parse_size(s.get("size", "小四"))
        para = {
            "line_spacing": s.get("line_spacing", 1.25),
            "space_before_pt": parse_length_pt(s.get("space_before", 0)),
            "space_after_pt": parse_length_pt(s.get("space_after", 0)),
            "align": s.get("align"),
            "first_line": s.get("first_line_indent"),
            "left": s.get("left_indent"),
            "hanging": s.get("hanging_indent"),
            "keep_next": s.get("keep_with_next"),
            "keep_lines": s.get("keep_lines"),
            "page_break_before": s.get("page_break_before"),
            "shading": s.get("shading"),
            "borders": s.get("borders"),
            "font_pt_for_chars": size_pt,
        }
        char = {
            "font_zh": s.get("font_zh"),
            "font_en": s.get("font_en"),
            "size_pt": size_pt,
            "bold": s.get("bold"),
            "italic": s.get("italic"),
            "color": s.get("color"),
            "underline": s.get("underline"),
        }
        return para, char

    def _effective_key(self, key: str) -> str:
        """引用块里的正文段落改用引用样式。"""
        depth = getattr(self, "_quote_depth", 0)
        if depth > 0 and key == "body":
            return "quote"
        return key

    def _new_paragraph(self, key: str, *, left_chars: float | None = None):
        key = self._effective_key(key)
        para, char = self._params(key)
        depth = getattr(self, "_quote_depth", 0)
        if key == "quote" and depth > 1:
            base_left = parse_chars(para.get("left") or "0字符") or 0.0
            para["left"] = f"{base_left + 2 * (depth - 1)}字符"
        p = self.doc.add_paragraph()
        p.style = self._ensure_style(STYLE_NAMES[key])
        if left_chars is not None:
            base_left = parse_chars(para.get("left") or "0字符") or 0.0
            para["left"] = f"{base_left + left_chars}字符"
        ooxml.apply_para_format(ooxml.para_ppr(p), **para)
        return p, char

    def _add_run(self, paragraph, text: str, ctx: RunCtx, char: dict):
        if not text:
            return None
        run = None
        if ctx.link and self.cfg["markdown"]["link"] == "hyperlink":
            r_id = self.doc.part.relate_to(ctx.link, ooxml.RT_HYPERLINK, is_external=True)
            hl = OxmlElement("w:hyperlink")
            hl.set(qn("r:id"), r_id)
            r_el = OxmlElement("w:r")
            hl.append(r_el)
            paragraph._p.append(hl)
            run = Run(r_el, paragraph)
        else:
            run = paragraph.add_run()
        run.text = text
        ooxml.preserve_space(run)

        fmt = dict(char)
        if ctx.mono:
            fmt["font_en"] = self.cfg["markdown"].get("code_font_en", "Consolas")
            fmt["font_zh"] = self.cfg["markdown"].get("code_font_zh", "宋体")
        if ctx.bold is not None:
            fmt["bold"] = ctx.bold
        if ctx.italic is not None:
            fmt["italic"] = ctx.italic
        if ctx.strike is not None:
            fmt["strike"] = ctx.strike
        if ctx.underline is not None:
            fmt["underline"] = ctx.underline
        if ctx.link and self.cfg["markdown"]["link"] in ("hyperlink", "text_with_url"):
            fmt["color"] = self.cfg["markdown"].get("link_color", "0563C1")
            fmt["underline"] = True
        ooxml.apply_run_format(ooxml.run_rpr(run), **fmt)
        return run

    # ------------------------------------------------------------------ 内联
    def _render_inline(self, tokens, paragraph, char: dict, base: RunCtx | None = None):
        md_cfg = self.cfg["markdown"]
        stack = [base or RunCtx()]
        last_char = ""
        for idx, tok in enumerate(tokens):
            t = tok.type
            ctx = stack[-1]
            if t == "text":
                self._add_run(paragraph, tok.content, ctx, char)
                last_char = tok.content[-1:]
            elif t == "strong_open":
                stack.append(replace(ctx, bold=md_cfg["strong_as_bold"]))
                continue
            elif t == "em_open":
                mark = md_cfg["emphasis_as_bold"]
                stack.append(replace(ctx, bold=True) if mark
                             else replace(ctx, italic=True))
                continue
            elif t == "s_open":
                stack.append(replace(ctx, strike=(md_cfg["delete"] == "strikethrough")))
                continue
            elif t == "link_open":
                stack.append(replace(ctx, link=tok.attrGet("href")))
                continue
            elif t in HEADING_CLOSE:
                if len(stack) > 1:
                    stack.pop()
                continue
            elif t == "code_inline":
                self._add_run(paragraph, tok.content, replace(ctx, mono=True), char)
                last_char = tok.content[-1:]
            elif t in ('math_inline', 'math_inline_double'):
                self._render_formula(tok.content, paragraph, char,
                                     display=t == 'math_inline_double')
                last_char = 'x'
            elif t == "softbreak":
                mode = md_cfg["softbreak"]
                nxt = self._next_text_char(tokens, idx)
                if mode == "newline":
                    paragraph.add_run().add_break()
                elif mode == "space":
                    self._add_run(paragraph, " ", ctx, char)
                elif mode == "auto":
                    if not (is_cjk(last_char) and (is_cjk(nxt) or nxt == "")):
                        self._add_run(paragraph, " ", ctx, char)
                continue
            elif t == "hardbreak":
                paragraph.add_run().add_break(WD_BREAK.LINE)
                continue
            elif t == "image":
                self._insert_inline_image(paragraph, tok)
                continue
            elif t == "html_inline":
                continue
            elif t in ("html_block", "emoji"):
                self._add_run(paragraph, tok.content, ctx, char)
            else:
                if tok.content and t not in HEADING_OPEN:
                    self._add_run(paragraph, tok.content, ctx, char)
                    last_char = tok.content[-1:]
            if ctx.link and md_cfg["link"] == "text_with_url":
                pass
        # link 文本后补 URL
        if md_cfg["link"] == "text_with_url":
            urls = [tk.attrGet("href") for tk in tokens if tk.type == "link_open"]
            for url in urls:
                if url:
                    self._add_run(paragraph, f"（{url}）", RunCtx(), char)

    @staticmethod
    def _next_text_char(tokens, idx: int) -> str:
        for tok in tokens[idx + 1:]:
            if tok.type == "text" and tok.content:
                return tok.content[0]
            if tok.type in ("softbreak", "hardbreak"):
                continue
            if tok.type in HEADING_OPEN:
                continue
            return ""
        return ""

    # ------------------------------------------------------------------ 图片
    def _render_formula(self, source, paragraph, char, display=False):
        mode = self.cfg.get('math', {}).get('mode', 'omml')
        size = float(char.get('size_pt') or parse_size(
            element_style(self.cfg, 'body').get('size', '小四')))
        try:
            if mode == 'text':
                self._add_run(paragraph, ('$$' if display else '$') + source +
                              ('$$' if display else '$'), RunCtx(), char)
                return
            if mode == 'omml':
                paragraph._p.append(latex_to_omml(source, display, size))
            else:
                image = render_png(source, display, size)
                width = image.width_pt * 12700
                height = image.height_pt * 12700
                scale = min(1.0, self.text_width_emu / width)
                run = paragraph.add_run()
                shape = run.add_picture(io.BytesIO(image.data),
                    width=Emu(round(width * scale)), height=Emu(round(height * scale)))
                # Preserve the source of our own PNGs for a later export to MD.
                shape._inline.docPr.set('descr', 'DocBridgeFormula:' + json.dumps(
                    {'latex': source, 'display': bool(display)}, ensure_ascii=False))
                shape._inline.docPr.set('title', '公式')
                if not display:
                    position = OxmlElement('w:position')
                    position.set(qn('w:val'), str(-round(image.depth_pt * scale * 2)))
                    run._r.get_or_add_rPr().append(position)
                self.counters['image'] += 1
            self.counters['formula'] += 1
        except FormulaError as exc:
            self.warnings.append(f'公式转换失败，已保留 LaTeX 源码：{exc}')
            self._add_run(paragraph, ('$$' if display else '$') + source +
                          ('$$' if display else '$'), RunCtx(), char)

    def _image_display_size(self, path: str, cfg: dict, natural_px=None):
        """按 96dpi 基准折算显示尺寸，再按文本栏宽/最大高度等比缩放。"""
        limit_w = int(self.text_width_emu * float(cfg.get("max_width_ratio", 1.0)))
        limit_h = None
        if cfg.get("max_height"):
            limit_h = int(parse_length_pt(cfg["max_height"]) / 72.0 * 914400)
        if natural_px:
            px_w, px_h = natural_px
        else:
            try:
                info = DocxImage.from_file(path)
                px_w, px_h = info.px_width, info.px_height
            except Exception:  # noqa: BLE001
                return None, None
        if not px_w or not px_h:
            return None, None
        w = int(px_w / 96.0 * 914400)
        h = int(px_h / 96.0 * 914400)
        scale = 1.0
        if w > limit_w:
            scale = limit_w / w
        if limit_h and h * scale > limit_h:
            scale = min(scale, limit_h / h)
        return int(w * scale), int(h * scale)

    def _insert_inline_image(self, paragraph, tok) -> None:
        src = (tok.attrGet("src") or "").strip()
        alt = ""
        if tok.children:
            alt = "".join(c.content for c in tok.children if c.type == "text")
        path = src
        if src and not re.match(r"^[a-zA-Z]+://", src) and not os.path.isabs(src):
            path = os.path.normpath(os.path.join(self.base_dir, unquote(src)))
        if not src or not os.path.exists(path):
            self.warnings.append(f"图片缺失，已跳过：{src or alt}")
            _p, char = self._params("caption")
            char = dict(char)
            char["color"] = "A6A6A6"      # 用灰字占位，不用斜体（全文不出现斜体）
            self._add_run(paragraph, f"[图片缺失：{alt or src}]", RunCtx(), char)
            return
        width, _height = self._image_display_size(path, element_style(self.cfg, "image"))
        run = paragraph.add_run()
        try:
            run.add_picture(path, width=Emu(width) if width else None)
        except Exception as exc:  # noqa: BLE001
            self.warnings.append(f"图片插入失败（{src}）：{exc}")
            return
        self.counters["image"] += 1

    # ------------------------------------------------------------------ Mermaid
    def _mermaid_renderer(self) -> MermaidRenderer:
        if self._mermaid is None:
            self._mermaid = MermaidRenderer(self.cfg.get("mermaid") or {},
                                            self._cache_dir, self.log)
        return self._mermaid

    def _render_mermaid_block(self, code: str, index: int) -> None:
        cfg = self.cfg["mermaid"]
        cfg = dict(cfg)
        cfg["_font_pt"] = parse_size(
            element_style(self.cfg, "code").get("size", "小四"))
        png = os.path.join(self._cache_dir, f"mermaid-{index:02d}.png")
        result = self._mermaid_renderer().render(code, png)
        if result is None:
            self.warnings.append(
                f"第 {index} 个 mermaid 图渲染失败，已按源码形式插入"
                "（检查 config 里的 mermaid.renderer / chrome_path）")
            if cfg.get("keep_source_on_fail", True):
                self._code_block_paragraphs(["```mermaid", *code.splitlines(), "```"])
            return
        self.counters["mermaid"] += 1
        para, _char = self._new_paragraph("image")
        para.alignment = ALIGN_ENUM.get(str(cfg.get("align", "center")).lower(),
                                        WD_ALIGN_PARAGRAPH.CENTER)
        width, height = self._image_display_size(
            result.path, cfg, natural_px=(result.width_px, result.height_px))
        para.add_run().add_picture(result.path, width=Emu(width), height=Emu(height))
        self.log(f"    · mermaid #{index} → {os.path.basename(result.path)} "
                 f"({result.width_px}×{result.height_px}px, {result.renderer})")

    # ------------------------------------------------------------------ 块级
    def render(self, text: str) -> None:
        # Generated docx2md anchors provide navigation in Markdown readers;
        # they are metadata, not Word body text.
        text = re.sub(r'^\s*<a id="[^"]+"></a>\s*$', '', text, flags=re.M)
        md_cfg = self.cfg["markdown"]
        if md_cfg.get("strip_front_matter") and text.lstrip().startswith("---"):
            text = re.sub(r"^\s*---\r?\n.*?\r?\n---\r?\n", "", text, count=1, flags=re.S)
        tokens = self.md.parse(text)
        if self.cfg["output"].get("toc"):
            self._insert_toc()
        self._walk(tokens, 0, len(tokens))

    def _insert_toc(self) -> None:
        title, char = self._new_paragraph("heading1")
        title_run = title.add_run("目录")
        ooxml.apply_run_format(ooxml.run_rpr(title_run), **char)
        para, body_char = self._new_paragraph("body")
        ooxml.add_toc_field(para, str(self.cfg["output"].get("toc_levels", "1-3")))
        for run in para.runs:      # 目录域的占位文字也要有字形，别用默认字体
            ooxml.apply_run_format(ooxml.run_rpr(run), **body_char)
        brk = self.doc.add_paragraph()
        brk.style = self._ensure_style("Normal")
        brk.add_run().add_break(WD_BREAK.PAGE)

    def _walk(self, tokens, start: int, end: int, list_stack=None) -> int:
        list_stack = list_stack or []
        i = start
        while i < end:
            tok = tokens[i]
            t = tok.type
            if t == 'math_block':
                para, char = self._new_paragraph('body')
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.first_line_indent = Emu(0)
                self._render_formula(tok.content.strip(), para, char, display=True)
                i += 1
                continue
            if t == "heading_open":
                level = int(tok.tag[1])
                key = f"heading{min(level, 6)}"
                inline = tokens[i + 1]
                para, char = self._new_paragraph(key)
                self._apply_heading_numbering(para, level)
                self._render_inline(
                    self._heading_tokens(inline.children or [], level), para, char)
                self.counters["heading"] += 1
                i += 3
                continue
            if t == "paragraph_open":
                inline = tokens[i + 1]
                hidden = bool(tok.hidden)
                if list_stack:
                    para, char = self._new_list_paragraph(list_stack[-1], hidden)
                else:
                    para, char = self._new_paragraph("body")
                self._render_inline(inline.children or [], para, char)
                self.counters["paragraph"] += 1
                i += 3
                continue
            if t in ("bullet_list_open", "ordered_list_open"):
                depth = len(list_stack) + 1
                num_id = BULLET_NUM_ID if t == "bullet_list_open" else ORDERED_NUM_ID
                list_stack.append({"num_id": num_id, "level": depth - 1})
                close = "bullet_list_close" if t == "bullet_list_open" else "ordered_list_close"
                j = i + 1
                depth_counter = 0
                while j < end:
                    if tokens[j].type == t:
                        depth_counter += 1
                    elif tokens[j].type == close:
                        if depth_counter == 0:
                            break
                        depth_counter -= 1
                    j += 1
                j = self._walk(tokens, i + 1, j, list_stack)
                list_stack.pop()
                self.counters["list"] += 1
                i = j + 1
                continue
            if t == "list_item_open":
                close = "list_item_close"
                j = i + 1
                depth_counter = 0
                while j < end:
                    if tokens[j].type == "list_item_open":
                        depth_counter += 1
                    elif tokens[j].type == close:
                        if depth_counter == 0:
                            break
                        depth_counter -= 1
                    j += 1
                self._item_first_paragraph = True
                j = self._walk(tokens, i + 1, j, list_stack)
                i = j + 1
                continue
            if t == "fence" or t == "code_block":
                info = (tok.info or "").strip().split()
                lang = info[0].lower() if info else ""
                if lang.startswith("mermaid"):
                    self.counters["code"] += 1
                    self._render_mermaid_block(tok.content.rstrip("\n"),
                                               self.counters["code"])
                else:
                    if lang == "text" and self.cfg["markdown"].get("warn_text_fence", True):
                        self.warnings.append(
                            "检测到 ```text 代码块：如果你是想画流程图/架构图，"
                            "建议改写成 ```mermaid")
                    self._code_block_paragraphs(tok.content.rstrip("\n").split("\n"))
                i += 1
                continue
            if t == "blockquote_open":
                close = "blockquote_close"
                depth = 1
                j = i + 1
                while j < end:
                    if tokens[j].type == "blockquote_open":
                        depth += 1
                    elif tokens[j].type == close:
                        depth -= 1
                        if depth == 0:
                            break
                    j += 1
                saved = list(list_stack)
                self._quote_depth = getattr(self, "_quote_depth", 0) + 1
                self._walk(tokens, i + 1, j, saved)
                self._quote_depth -= 1
                i = j + 1
                continue
            if t == "table_open":
                close = "table_close"
                j = i + 1
                while j < end and tokens[j].type != close:
                    j += 1
                self._render_table(tokens, i, j)
                self.counters["table"] += 1
                i = j + 1
                continue
            if t == "hr":
                _para, _char = self._new_paragraph("hr")
                i += 1
                continue
            if t == "html_block":
                self.warnings.append("已忽略一段内嵌 HTML（Markdown 里不建议写裸 HTML）")
                i += 1
                continue
            if t in ("thead_open", "tbody_open", "tr_open", "th_open", "td_open",
                     "bullet_list_close", "ordered_list_close", "list_item_close"):
                i += 1
                continue
            if t == "inline":
                para, char = self._new_paragraph("body")
                self._render_inline(tok.children or [], para, char)
                i += 1
                continue
            i += 1
        return i

    # -------------------------------------------------------------- 标题编号
    def _numbering_level(self, level: int) -> int | None:
        """Markdown 标题层级 → 编号定义里的 ilvl；不参与编号时返回 None。

        编号定义里 ilvl=0 就是"编号的第 1 级"（见 `numbering.heading_plan`）。
        不跳级时 Markdown 第 N 级标题接 ilvl=N-1。
        开了 `skip_first_level` 后整体下移一位：第 1 级标题是整篇的主标题、
        不编号，第 2 级标题才接编号第 1 级，即 ilvl=N-2。

        `levels: n` 表示编号定义里有 n 级可用（ilvl 0..n-1），超出的一律不编号
        ——这时它也带不出上级序号，与 Word 自身的 multilevel 行为一致。
        """
        cfg = self._heading_numbering
        if not cfg:
            return None
        skip = bool(cfg.get("skip_first_level", False))
        if skip and level <= 1:
            return None                      # 主标题不编号
        ilvl = level - (2 if skip else 1)
        if ilvl >= int(cfg.get("levels", 6)):
            return None
        return ilvl

    def _apply_heading_numbering(self, paragraph, level: int) -> None:
        """把标题接进 Word 的多级列表编号体系。

        写的是真正的 `w:numPr`（列表编号），不是正文里的文本编号——所以在
        Word 里增删章节时序号会自动重排，交叉引用也能用。

        层级与编号定义的对齐靠 `numbering.preset`：预设的第 N 项对应第 N 级
        标题，超出 `levels` 的深层标题不编号（这时它也带不出上级序号，与
        Word 自身的 multilevel 行为一致）。开了 `skip_first_level` 时，
        第 1 级标题保持 Word 标题样式但不挂编号。具体映射见 `_numbering_level`。
        """
        cfg = self._heading_numbering
        if not cfg:
            return
        ilvl = self._numbering_level(level)
        if ilvl is None:
            return
        para_cfg = {
            "num_id": HEADING_NUM_ID,
            "num_level": ilvl,
        }
        left = parse_chars(cfg.get("indent") or "0字符")
        hanging = parse_chars(cfg.get("hanging") or "0字符")
        if left:
            para_cfg["left"] = f"{left}字符"
        if hanging and hanging > 0:
            para_cfg["hanging"] = f"{hanging}字符"
        ooxml.apply_para_format(ooxml.para_ppr(paragraph), **para_cfg)

    def _heading_tokens(self, tokens, level: int):
        """如果标题手写了编号前缀，且已开自动编号，就把前缀从文字里去掉。

        只在**第一个**文本 token 上剥，且剥完为空时保留原样——避免把标题
        剥没了。`# 一、项目概述` + 自动编号 → "一、项目概述"，不处理就是
        "一、一、项目概述"。
        """
        cfg = self._heading_numbering
        if not cfg or not cfg.get("strip_text_prefix", True):
            return tokens
        out = list(tokens)
        for index, tok in enumerate(out):
            if tok.type == "text" and tok.content and tok.content.strip():
                cleaned = strip_heading_prefix(tok.content)
                if cleaned != tok.content:
                    new_tok = tok.copy()
                    new_tok.content = cleaned
                    out[index] = new_tok
                return out
            if tok.type in ("code_inline", "image"):
                return out          # 标题以行内代码/图片开头，不剥
        return out

    def _new_list_paragraph(self, frame: dict, tight: bool):
        key = "list" if frame["num_id"] == BULLET_NUM_ID else "list_ordered"
        para, char = self._new_paragraph(key)
        s = element_style(self.cfg, key)
        size_pt = parse_size(s.get("size", "小四"))
        base_left = parse_chars(s.get("left_indent") or "2字符") or 2.0
        step = float(s.get("level_step_chars", 2))
        para_cfg = {"line_spacing": s.get("line_spacing", 1.25),
                    "space_before_pt": 0 if tight else parse_length_pt(s.get("space_before", 0)),
                    "space_after_pt": 0 if tight else parse_length_pt(s.get("space_after", 0)),
                    "align": s.get("align"),
                    "left": f"{base_left + step * frame['level']}字符",
                    "hanging": s.get("hanging_indent", "2字符"),
                    "first_line": "0字符",
                    "font_pt_for_chars": size_pt}
        first = getattr(self, "_item_first_paragraph", True)
        if first:
            para_cfg["num_id"] = frame["num_id"]
            para_cfg["num_level"] = frame["level"]
        else:
            para_cfg["left"] = f"{base_left + step * frame['level'] + 2}字符"
            para_cfg["hanging"] = "0字符"
        ooxml.apply_para_format(ooxml.para_ppr(para), **para_cfg)
        self._item_first_paragraph = False
        return para, char

    def _quote_depth_para(self, key: str, extra_chars: float):
        para, char = self._new_paragraph(key, left_chars=extra_chars)
        return para, char

    def _code_block_paragraphs(self, lines: list[str]) -> None:
        """代码块：一行一段，段间无间距 + 底纹，视觉上连成一个色块。"""
        for line in lines or [""]:
            para, char = self._new_paragraph("code")
            run = para.add_run(line if line else "\u00a0")
            ooxml.preserve_space(run)
            fmt = dict(char)
            fmt["font_en"] = self.cfg["markdown"].get("code_font_en", "Consolas")
            fmt["font_zh"] = self.cfg["markdown"].get("code_font_zh", "宋体")
            ooxml.apply_run_format(ooxml.run_rpr(run), **fmt)
        self.counters["code"] += 1

    def _render_table(self, tokens, start: int, end: int) -> None:
        rows: list[tuple[bool, list]] = []
        aligns: list[str] = []
        i = start
        current: list | None = None
        is_header = False
        header_rows: set[int] = set()
        while i < end:
            t = tokens[i].type
            if t == "thead_open":
                is_header = True
            elif t == "thead_close":
                is_header = False
            elif t == "tr_open":
                current = []
            elif t == "tr_close":
                if current is not None:
                    rows.append((len(rows) in header_rows or is_header, current))
                current = None
            elif t == "th_open" or t == "td_open":
                style = tokens[i].attrGet("style") or ""
                m = re.search(r"text-align:\s*(left|center|right)", style)
                if m and len(aligns) == len(current or []):
                    aligns.append(m.group(1))
                elif len(aligns) == len(current or []):
                    aligns.append("")
                if t == "th_open":
                    header_rows.add(len(rows))
                i += 1
                if i < end and tokens[i].type == "inline":
                    current.append(tokens[i].children or [])
                i += 1
                continue
            i += 1

        if not rows:
            return
        n_cols = max(len(r[1]) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=n_cols)
        ooxml.set_table_borders(table, self.cfg.get("table", {}).get("border"))
        ooxml.set_table_width_pct(table, self.cfg.get("table", {}).get("width_pct", 100))
        ooxml.set_table_layout_fixed(table, bool(self.cfg.get("table", {}).get("fixed_layout", False)))
        margins = self.cfg.get("table", {}).get("cell_margin") or {}
        ooxml.set_table_cell_margins(table, **margins) if margins else ooxml.set_table_cell_margins(table)

        for r, (header, cells) in enumerate(rows):
            key = "table_header" if header else "table_text"
            if header:
                ooxml.repeat_header_row(table.rows[r])
            for c in range(n_cols):
                cell = table.cell(r, c)
                ooxml.set_cell_valign(cell, self.cfg.get("table", {}).get("valign", "center"))
                if header and self.cfg.get("table", {}).get("header_fill"):
                    ooxml.set_cell_shading(cell, self.cfg["table"]["header_fill"])
                token_list = cells[c] if c < len(cells) else []
                cell.text = ""
                para = cell.paragraphs[0]
                para.style = self._ensure_style(STYLE_NAMES[key])
                s = element_style(self.cfg, key)
                custom = aligns[c] if c < len(aligns) and aligns[c] else s.get("align")
                ooxml.apply_para_format(
                    ooxml.para_ppr(para),
                    line_spacing=s.get("line_spacing", 1.25),
                    space_before_pt=parse_length_pt(s.get("space_before", 0)),
                    space_after_pt=parse_length_pt(s.get("space_after", 0)),
                    align=custom, first_line="0字符",
                    font_pt_for_chars=parse_size(s.get("size", "五号")))
                _p, char = self._params(key)
                self._render_inline(token_list, para, char)

    # ------------------------------------------------------------------ 出口
    def _numbering_plans(self) -> list:
        """组装要写进 numbering.xml 的全部编号定义。"""
        list_key = element_style(self.cfg, "list")
        ordered_key = element_style(self.cfg, "list_ordered")
        list_cfg = {
            **(self.cfg.get("elements", {}).get("list") or {}),
            "_font_pt": parse_size(list_key.get("size", "小四")),
            "hanging_indent_chars": parse_chars(
                list_key.get("hanging_indent") or "2字符") or 2.0,
        }
        ordered_cfg = {
            **(self.cfg.get("elements", {}).get("list_ordered") or {}),
            "_font_pt": parse_size(ordered_key.get("size", "小四")),
            "hanging_indent_chars": parse_chars(
                ordered_key.get("hanging_indent") or "2字符") or 2.0,
        }
        plans = list_plans(list_cfg, ordered_cfg)
        if self._heading_numbering:
            font_pt = parse_size(
                element_style(self.cfg, "heading1").get("size", "小三"))
            plans.append(heading_plan(self._heading_numbering, font_pt))
        return plans

    def save(self) -> None:
        props = self.cfg.get("output") or {}
        core = self.doc.core_properties
        if props.get("title"):
            core.title = str(props["title"])
        if props.get("author"):
            core.author = str(props["author"])
        ensure_numbering_part(self.doc, self._numbering_plans())
        os.makedirs(os.path.dirname(self.out_path) or ".", exist_ok=True)
        self.doc.save(self.out_path)
        if self._mermaid:
            self._mermaid.close()
